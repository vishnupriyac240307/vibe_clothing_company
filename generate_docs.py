import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_table_header(row, col_widths, headers, bg_hex="0F172A"):
    for idx, header_text in enumerate(headers):
        cell = row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, bg_hex)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header_text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)

def build_docx_documentation():
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Colors
    NAVY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    GRAY = RGBColor(100, 116, 139)

    # Title Page
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(36)

    # Add Company Logo before Company Name
    logo_path = os.path.abspath(os.path.join(os.path.dirname(__file__), 'static', 'images', 'logo.png'))
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(18)
        p_logo.paragraph_format.space_after = Pt(12)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.2))

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("VIBE CLOTHING COMPANY\nSMART PRODUCTION & MANUFACTURING MANAGEMENT SYSTEM")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("PHASE 2 PROJECT DOCUMENTATION")
    r_sub.bold = True
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = BLUE

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(108)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("A Real-Time Clothing Manufacturing & Production Management System\nBuilt with Python, Flask, SQLAlchemy, MySQL/SQLite & Chart.js\n\nAcademic Year 2025 – 2026")
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = GRAY

    doc.add_page_break()

    # Helper function for headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = NAVY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = BLUE
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = NAVY
        return h

    def add_p(text):
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        return p

    # ACKNOWLEDGEMENT
    add_h1("ACKNOWLEDGEMENT")
    add_p("I express my sincere gratitude and deep appreciation to everyone who supported the successful design, development, testing, and implementation of the Vibe Clothing Company Smart Production & Manufacturing Management System (ERP).")
    add_p("I extend my heartfelt thanks to the faculty members, technical advisors, and industry experts for their invaluable guidance, constructive feedback, and continuous encouragement throughout the Phase 2 implementation of this project.")
    add_p("Special thanks to Vibe Clothing Company operational stakeholders for providing domain insights into garment manufacturing, Bill of Materials (BOM) management, multi-stage batch tracking, quality control standards, and outbound logistics workflows.")

    # CONTENTS
    add_h1("CONTENTS")
    toc_items = [
        ("SYNOPSIS", "iii"),
        ("A. INTRODUCTION", "1"),
        ("   1. Organization Profile", "1"),
        ("   2. System Specification", "2"),
        ("   3. Hardware Configuration", "3"),
        ("   4. Software Specification", "4"),
        ("B. SYSTEM STUDY", "5"),
        ("   1. Existing System", "5"),
        ("   2. Drawbacks", "6"),
        ("   3. Proposed System", "7"),
        ("   4. Features", "8"),
        ("C. SYSTEM DESIGN AND DEVELOPMENT", "10"),
        ("   1. File Design", "10"),
        ("   2. Input Design", "12"),
        ("   3. Output Design", "15"),
        ("   4. Database Design", "18"),
        ("   5. System Development", "22"),
        ("   6. Description of Modules", "25"),
        ("4. SOFTWARE TESTING AND IMPLEMENTATION", "42"),
        ("CONCLUSION", "47"),
        ("BIBLIOGRAPHY", "48"),
        ("APPENDICES", "49"),
        ("   D. Data Flow Diagram", "49"),
        ("   E. Table Structure", "52"),
        ("   F. Sample Coding", "58"),
        ("   G. Sample Input", "65"),
        ("   H. Sample Output", "68")
    ]
    for topic, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(topic)
        r1.bold = not topic.startswith("   ")
        r2 = p.add_run(f" ........................................................................................................ {page}")
        r2.font.color.rgb = GRAY

    # SYNOPSIS
    add_h1("SYNOPSIS")
    add_p("The Vibe Clothing Company Smart Production & Manufacturing Management System is a commercial-grade, real-time Enterprise Resource Planning (ERP) platform developed specifically for modern garment production facilities. Traditional clothing manufacturing relies heavily on disconnected spreadsheets, manual logbooks, and isolated inventory channels, leading to raw material stockouts, production bottlenecks, poor quality tracking, delayed shipments, and an absence of unified cost and operational visibility.")
    add_p("This application solves these operational challenges by implementing a single, fully normalized relational database engine (MySQL/SQLite) coupled with a Python Flask REST backend, server-side Role-Based Access Control (RBAC), an automated BOM & production engine (ERPEngine), interactive Chart.js BI analytics, visual 5-stage garment batch tracking, quality control pass-rate analytics, outbound shipment dispatches, and multi-format reporting (PDF/Excel/CSV).")
    add_p("All financial values across the entire platform are strictly processed and displayed in Indian Rupees (₹) with standard Indian formatting (e.g. ₹12,50,000.00). Furthermore, products are identified strictly by unique Product IDs (PROD-0001), completely eliminating outdated SKU terminology. Automated triggers and business logic ensure that raw materials consumed during production update inventory ledgers instantly, QC-passed units directly dictate finished goods stock, and customer shipments deduct finished goods stock while tracking delivery milestones.")

    doc.add_page_break()

    # A. INTRODUCTION
    add_h1("A. INTRODUCTION")

    add_h2("1. Organization Profile")
    add_p("Vibe Clothing Company is a modern apparel manufacturing enterprise specializing in high-quality knitted and woven garments including Cotton T-Shirts, Slim-Fit Denim Jeans, Heavyweight Denim Trucker Jackets, and Fleece-Lined Pullover Hoodies. The organization manages end-to-end manufacturing workflows comprising raw material procurement (combed cotton fabric, merino wool, high-tenacity thread, resin buttons, eco-friendly reactive dyes), garment cutting, dyeing, stitching, finishing, quality inspection, packaging, and wholesale B2B dispatch to major retail and e-commerce clients such as Flipkart, Amazon India, and Nykaa.")
    add_p("Prior to the deployment of this Smart ERP system, plant operations suffered from fragmented tracking between warehouse stock, floor supervisor logs, QA defect registers, and logistics dispatch slips. The implementation of this integrated ERP system establishes total operational visibility, traceability from customer order down to individual fabric lot, automated Bill of Materials (BOM) stock reservation, and real-time executive dashboard decision-making.")

    add_h2("2. System Specification")
    add_p("The Vibe Clothing Company ERP is built as a modular client-server web application adhering to the Model-View-Controller (MVC) and RESTful API architecture:")
    add_bullet("Architecture: Web-Based Client-Server Model-View-Controller (MVC)")
    add_bullet("Backend Engine: Python 3.14 with Flask 3.0 Web Framework")
    add_bullet("ORM Layer: Flask-SQLAlchemy 3.1 with PyMySQL driver for MySQL / SQLite DB Engine")
    add_bullet("Frontend Stack: HTML5, CSS3 (Custom Dark Navy & Slate Theme), JavaScript (ES6+), Bootstrap 5.3.2, FontAwesome 6 Icons")
    add_bullet("Data Visualization: Chart.js 4.4.1 for real-time BI analytics and factory trend reporting")
    add_bullet("Reporting Engines: ReportLab 4.0 for programmatic PDF generation; openpyxl 3.1 for Excel spreadsheet exports")
    add_bullet("Authentication & Security: Werkzeug scrypt password hashing, HTTP session management, server-side RBAC decorators (@role_required)")

    add_h2("3. Hardware Configuration")
    add_p("The system is optimized to operate efficiently on standard off-the-shelf development, server, and client hardware without requiring proprietary industrial hardware:")
    
    add_h3("Development & Production Server Requirements:")
    add_bullet("Processor: Quad-Core Intel Core i5 / Apple Silicon M1 or equivalent (2.0 GHz+)")
    add_bullet("RAM: 8 GB DDR4 / Unified Memory (16 GB recommended for high concurrent loads)")
    add_bullet("Storage: 256 GB SSD (Solid State Drive) with minimum 10 GB free space for DB & media assets")
    add_bullet("Network: High-speed Ethernet / 802.11ac Wi-Fi interface for REST API communication")

    add_h3("Client Workstation / Mobile Requirements:")
    add_bullet("Processor: Dual-Core 1.5 GHz or higher")
    add_bullet("RAM: 2 GB (Mobile) / 4 GB (Desktop Workstation)")
    add_bullet("Display Resolution: Responsive interface supporting 320px (Mobile), 768px (Tablet), 1366x768 / 1920x1080 (Desktop)")
    add_bullet("Web Browser: Modern HTML5 browser (Google Chrome 110+, Apple Safari 16+, Mozilla Firefox 110+, Microsoft Edge)")

    add_h2("4. Software Specification")
    add_p("The software components and third-party libraries utilized in the application include:")
    add_bullet("Operating System: macOS / Linux (Ubuntu 22.04 LTS) / Microsoft Windows 11")
    add_bullet("Programming Language: Python 3.14.2")
    add_bullet("Web Framework: Flask 3.0.0")
    add_bullet("Database Management System: SQLite 3 (Default Local Instance) / MySQL 8.0 Enterprise Server")
    add_bullet("Database Driver: PyMySQL 1.1.0")
    add_bullet("PDF Generation Library: ReportLab 4.0.0")
    add_bullet("Excel Generation Library: openpyxl 3.1.2")
    add_bullet("Security & Utility Libraries: Werkzeug 3.0.0, python-dotenv 1.0.0, requests 2.31.0")

    # B. SYSTEM STUDY
    add_h1("B. SYSTEM STUDY")

    add_h2("1. Existing System")
    add_p("In the traditional manufacturing setup at Vibe Clothing Company, operational departments functioned in silos:")
    add_bullet("Warehouse staff recorded fabric and trim stock manually on paper registers or local Excel sheets.")
    add_bullet("Production planning managers created cutting orders without real-time verification of raw material readiness.")
    add_bullet("Floor supervisors tracked garment assembly stages (Cutting, Dyeing, Stitching, Finishing, Packing) on whiteboard charts.")
    add_bullet("Quality inspection results were written on manual QA slips, making defect rate analysis cumbersome.")
    add_bullet("Logistics dispatches were handled independently without automated stock deduction or direct customer order linking.")

    add_h2("2. Drawbacks of the Existing System")
    add_bullet("Frequent Material Shortages: Production started without verifying full BOM availability, leading to half-stitched garments sitting idle on shop floors.")
    add_bullet("Lack of Traceability: Inability to trace a delivered customer batch back to the exact fabric lot or QA inspector.")
    add_bullet("Manual Calculation Errors: Human errors in valuation, pass-rate calculations, and unit conversions.")
    add_bullet("Delayed Decision Making: Management received paper reports days or weeks after production completed.")
    add_bullet("Security & Audit Deficits: Unrestricted access to local spreadsheets without audit logs or role enforcement.")

    add_h2("3. Proposed System")
    add_p("The Vibe Clothing Company Smart Production & Manufacturing Management System establishes an integrated digital manufacturing platform:")
    add_bullet("Centralized Relational Database: Single source of truth connecting all 13 core database models.")
    add_bullet("Automated BOM & Readiness Scoring: Calculates exact raw material readiness (READY, PARTIALLY_READY, BLOCKED) before starting production.")
    add_bullet("Automated Stock Ledger Transactions: Consumes raw stock on production start and adds usable finished goods on order completion.")
    add_bullet("QC Pass-Rate Integration: Only QA-passed garments (Total - Defective) enter usable finished goods stock.")
    add_bullet("Real-Time BI Dashboard: 10 dynamic KPI cards, Chart.js trend visuals, global search (Ctrl K), and live notifications drawer.")

    add_h2("4. Key Features Implemented")
    add_bullet("Indian Rupee (₹) Currency Standard: All financial metrics, inventory valuations, cost prices, and sales revenues use ₹ INR with standard Indian formatting.")
    add_bullet("Product ID Standard: Strictly uses Product ID (PROD-0001). SKU terminology is completely eliminated.")
    add_bullet("5-Stage Visual Batch Stepper: Interactive stage advancement across Cutting → Dyeing → Stitching → Finishing → Packing.")
    add_bullet("Role-Based Access Control: Strict server-side RBAC enforcing Admin, Manager, and Staff permission levels.")
    add_bullet("Multi-Format Report Exporter: Programmatic PDF, Excel (.xlsx), and CSV exports for all main modules.")
    add_bullet("Full End-to-End Traceability: Lineage graph mapping Customer → Order → Product → BOM → Batch → QC → Shipment.")

    # C. SYSTEM DESIGN AND DEVELOPMENT
    add_h1("C. SYSTEM DESIGN AND DEVELOPMENT")

    add_h2("1. File Design & Directory Structure")
    add_p("The project directory architecture follows a clean, modular Flask organization:")
    add_bullet("app.py: Core application factory, blueprint registrations, context processors, error handlers, and auto-seeder.")
    add_bullet("config.py: Environment configuration, database URI routing, secret key handling.")
    add_bullet("db.py: SQLAlchemy database instance initialization.")
    add_bullet("models/: 13 Relational entity models (User, Material, Product, ProductBOM, Customer, Supplier, ProductionOrder, Batch, QualityInspection, Shipment, InventoryTransaction, ActivityLog, Notification).")
    add_bullet("routes/: 17 Blueprint API controllers handling HTTP requests, form submissions, and view rendering.")
    add_bullet("services/: Business logic engines (ERPEngine, ImageService, LogoService, ExportService).")
    add_bullet("utils/: IdGenerator (PROD-XXXX, MAT-XXXX auto-formatting), Helpers (format_inr, log_activity, create_notification), Decorators (login_required, role_required).")
    add_bullet("templates/: 23 Jinja2 HTML templates covering all views, modals, and error pages.")
    add_bullet("static/: Assets including custom.css, app.js, charts.js, real-time.js, placeholder.svg, logo.png, and login-background.mp4.")

    add_h2("2. Input Design")
    add_p("Input forms in the application feature strict client-side HTML5 constraints and server-side validation:")
    add_bullet("Sign In Form (login.html): Captures username/email and password with session remember checkbox and password visibility toggle.")
    add_bullet("Add Raw Material Form (inventory/index.html): Captures Material Name, Code (MAT-XXX), Category, Unit, Unit Cost (₹), Minimum Level, Reorder Level, Critical Level, and Supplier.")
    add_bullet("Add Product Form (products/index.html): Captures Product Name, Category, Size, Color, Selling Price (₹), Cost Price (₹), Target Quantity, and Image URL.")
    add_bullet("Bill of Materials Form (products/index.html): Links Product ID to Raw Material IDs with quantity required per unit.")
    add_bullet("Production Order Form (production/index.html): Captures Product, Target Quantity, Customer, Priority, Expected Completion Date, and Supervisor.")
    add_bullet("QC Inspection Form (quality/index.html): Captures Batch ID, Total Evaluated Garments, Defective Units, Defect Category, and QA Remarks.")
    add_bullet("Shipment Form (shipments/index.html): Captures Customer, Product, Batch, Quantity, Carrier, Tracking Number, and Expected Delivery Date.")

    add_h2("3. Output Design")
    add_p("The system produces clean, structured outputs formatted for industrial usability:")
    add_bullet("Dashboard Overview: 10 KPI summary cards, active batch list, low stock warning table, and order distribution donut chart.")
    add_bullet("Inventory Stock Ledger: Tabular material master with live status badges (OK, LOW, CRITICAL) and total valuation in ₹.")
    add_bullet("Kanban Production Board: 4-column drag/click workflow cards (PLANNED, IN PROGRESS, COMPLETED, CANCELLED) with readiness percentage pills.")
    add_bullet("Batch Stage Stepper: Visual progress bar displaying active stage and produced vs target garment counts.")
    add_bullet("Quality Control Register: Pass/Fail status badges, defect metrics, pass rate percentage progress bars.")
    add_bullet("Factory Analytics BI: 6 Chart.js graphs displaying Output Timeline, Order Breakdown, Top Products Stock, Inventory Category Valuation in ₹, Defect Categories, and Customer Lifetime Spend in ₹.")
    add_bullet("Export Files: Formatted PDF invoices/reports, Excel worksheets with bold headers and ₹ currency formatting, and CSV data dumps.")

    add_h2("4. Database Design")
    add_p("The application utilizes a normalized relational schema with 13 entities, strict foreign keys, index optimization, and cascade rules:")

    # Table of Entities
    t_db = doc.add_table(rows=1, cols=4)
    t_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_header(t_db.rows[0], [Inches(1.5), Inches(1.5), Inches(1.2), Inches(2.3)], ["Table Name", "Primary Key", "Engine", "Description"])
    
    db_tables_meta = [
        ("users", "id / user_id", "InnoDB", "System users, roles, password hashes, and login history"),
        ("materials", "id / material_id", "InnoDB", "Raw materials, fabric rolls, stock levels, unit costs in ₹"),
        ("products", "id / product_id", "InnoDB", "Finished garment catalog, selling/cost prices in ₹, stock"),
        ("product_bom", "id", "InnoDB", "Bill of Materials mapping product to raw materials & unit requirements"),
        ("customers", "id / customer_id", "InnoDB", "B2B client profiles, company names, logos, GST IDs"),
        ("suppliers", "id / supplier_id", "InnoDB", "Vendor profiles, categories, payment terms, ratings"),
        ("production_orders", "id / order_id", "InnoDB", "Production manufacturing orders, target quantities, readiness scores"),
        ("batches", "id / batch_id", "InnoDB", "Garment batches, 5-stage stepper progress, produced quantities"),
        ("quality_inspections", "id / inspection_id", "InnoDB", "QA inspection records, total evaluated, defects, pass rates"),
        ("shipments", "id / shipment_id", "InnoDB", "Outbound dispatches, tracking numbers, carriers, delivery status"),
        ("inventory_transactions", "id / transaction_id", "InnoDB", "Stock movements (Stock In, Production Consumption, Adjustment)"),
        ("activity_logs", "id", "InnoDB", "System audit trail logging user actions and record IDs"),
        ("notifications", "id", "InnoDB", "System alerts, low stock warnings, QC failures, delivery alerts")
    ]
    for tab, pk, eng, desc in db_tables_meta:
        row = t_db.add_row()
        row.cells[0].paragraphs[0].text = tab
        row.cells[1].paragraphs[0].text = pk
        row.cells[2].paragraphs[0].text = eng
        row.cells[3].paragraphs[0].text = desc
        for c in row.cells:
            set_cell_margins(c)

    add_h2("5. System Development & Architecture")
    add_p("The system development followed modular software engineering practices:")
    add_bullet("ERPEngine Business Service: Handles complex transactional logic in Python. When start_production() is called, ERPEngine validates material availability, deducts required fabric/trims from Material.quantity, creates an InventoryTransaction log, generates a Batch, and updates ProductionOrder status to IN_PROGRESS.")
    add_bullet("Quality & Finished Goods Linkage: When a QualityInspection is recorded, passed_units = total_evaluated - defective_units is computed. When complete_production() runs, product.available_stock is incremented by passed_units (not defective units).")
    add_bullet("Shipment Finished Goods Deduction: When a Shipment is added, Product.available_stock is validated and reduced by the shipment quantity.")
    add_bullet("AJAX Polling Real-Time Feed: static/js/real-time.js polls /api/notifications every 15 seconds to update the notification bell drawer and unread counter badge dynamically.")

    add_h2("6. Description of Implemented Modules")

    modules_info = [
        ("1. Login and Authentication Module", "Handles secure user authentication via username/email and password using Werkzeug scrypt hashing. Features remember me session handling, password visibility toggle, full-screen video background (login-background.mp4), dark overlay, and demo credentials modal."),
        ("2. Signup / User Registration Module", "Allows new users to register an account with validation for duplicate username/email, minimum password length (6+ chars), and role assignment."),
        ("3. Role-Based Access Control (RBAC) Module", "Enforces server-side authorization decorators (@role_required) across routes. Admin gets full access including user management and activity logs; Manager gets operational access; Staff gets entry/view access."),
        ("4. Dashboard (Production Overview) Module", "Displays 10 live KPI summary cards (Total Materials, Products, Active Orders, Completed Orders, Active Batches, Low Stock Items, Efficiency %, Quality Pass Rate %, Shipments in Transit, Inventory Valuation in ₹), active batch progress, low stock alerts, and order distribution chart."),
        ("5. Inventory Management Module", "Provides raw material master table, unit cost in ₹, total value calculation, stock status badges (OK, LOW, CRITICAL), add material modal, and record stock movement modal."),
        ("6. Material Management Module", "Handles raw fabric/trim creation, reorder thresholds, storage locations (Rack A-12, Bin T-04), supplier links, and lot/batch tracking numbers."),
        ("7. Product Management & BOM Module", "Manages finished garment catalog (PROD-0001), sizes, colors, selling price in ₹, cost price in ₹, image lookups, and Bill of Materials (BOM) ingredient specifications."),
        ("8. Production Management (Kanban) Module", "Displays 4-column Kanban board (PLANNED, IN PROGRESS, COMPLETED, CANCELLED) with readiness score percentage calculation and Start/Complete Order actions."),
        ("9. Customer Management Module", "Tracks B2B customer profiles (Flipkart, Amazon India, Nykaa), Clearbit/UI-Avatars company logos, GST IDs, order histories, and lifetime spend in ₹."),
        ("10. Supplier Management Module", "Manages vendor records (ABC Textiles, Premium Threads), supplier categories, payment terms (Net 30, Advance), vendor ratings out of 5.0, and supplied materials list."),
        ("11. Batch Management & 5-Stage Stepper Module", "Tracks garment manufacturing batches (BAT-0001) through 5 stages: Cutting → Dyeing → Stitching → Finishing → Packing with interactive stage advancement buttons."),
        ("12. Quality Control (QC) Module", "Logs garment quality inspections (QC-0001), total evaluated, defective units, defect types (Minor Loose Stitching), calculates Pass Rate % and flags PASS/FAIL status."),
        ("13. Shipment Management & Logistics Module", "Manages outbound shipment dispatches (SHP-0001), carrier assignment (Blue Dart Logistics), tracking numbers, expected delivery dates, and status updates to DELIVERED."),
        ("14. Traceability Lineage Graph Module", "Visual node-based lineage timeline tracing Customer → Production Order → Product → BOM/Materials → Batch Stepper → Quality Control → Shipment Delivery."),
        ("15. Analytics & BI Module", "Renders 6 interactive Chart.js visualizations: Output Timeline, Order Distribution, Top Products Stock, Inventory Category Valuation in ₹, Defect Categories, and Customer Spend in ₹ with date range filtering (Today, 7D, 30D, 3M)."),
        ("16. User Management Directory Module", "Admin-only user directory displaying user cards, active/inactive archive tabs, create user modal, and role editing permissions."),
        ("17. Activity Audit Logs Module", "Provides an immutable audit log trail capturing user action, target module, record ID, description, timestamp, and IP address for compliance."),
        ("18. System Notifications Module", "Manages system notification drawer alerts (danger, warning, info, success), unread counter badge, and mark all read API endpoint."),
        ("19. Global Search (Ctrl K) Module", "Renders interactive search modal searching instantly across Materials, Products, Orders, Batches, Customers, Suppliers, and Shipments."),
        ("20. Reports & Exports Engine Module", "Generates ReportLab PDF invoices/reports, openpyxl Excel spreadsheets with bold headers and ₹ currency formatting, and raw CSV files across all modules.")
    ]

    for title, desc in modules_info:
        add_h3(title)
        add_p(desc)

    # 4. SOFTWARE TESTING AND IMPLEMENTATION
    add_h1("4. SOFTWARE TESTING AND IMPLEMENTATION")
    add_p("Comprehensive automated and manual testing was performed across all 20 modules to verify system integrity, data consistency, security, and performance.")

    add_h2("Automated & Integration Test Execution")
    add_p("An automated test suite (test_app.py) was built using Python unittest framework. Execution results:")
    add_bullet("test_01_currency_formatting: PASSED (Verified ₹ formatting without USD)")
    add_bullet("test_02_id_generator: PASSED (Verified PROD-XXXX, MAT-XXXX auto-IDs)")
    add_bullet("test_03_authentication_and_rbac: PASSED (Verified Admin access vs Staff 403 Forbidden)")
    add_bullet("test_04_product_no_sku: PASSED (Verified Product ID and NO SKU attribute)")
    add_bullet("test_05_material_availability_and_readiness: PASSED (Verified BOM calculations)")
    add_bullet("test_06_end_to_end_production_workflow: PASSED (Verified full 21-step workflow)")

    add_h2("System Test Cases Table")
    t_test = doc.add_table(rows=1, cols=6)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_header(t_test.rows[0], [Inches(0.8), Inches(1.2), Inches(1.5), Inches(1.2), Inches(1.2), Inches(0.6)], ["TC ID", "Module", "Scenario", "Input", "Expected Result", "Status"])

    test_cases_data = [
        ("TC-01", "Auth", "Login with valid credentials", "admin / admin123", "Login success, redirect to dashboard", "PASS"),
        ("TC-02", "RBAC", "Staff access user management", "GET /users as staff", "403 Forbidden page rendered", "PASS"),
        ("TC-03", "BOM", "Calculate readiness for 100 T-Shirts", "PROD-0001, Qty 100", "Availability 100%, Status READY", "PASS"),
        ("TC-04", "Production", "Start order with sufficient materials", "ORD-0001 start", "Raw stock deducted, batch created", "PASS"),
        ("TC-05", "Production", "Start order with insufficient materials", "ORD-0002 start (blocked)", "Blocked error alert, stock untouched", "PASS"),
        ("TC-06", "Quality", "Record QC inspection with 5 defects", "Total 100, Defects 5", "Pass rate 95%, Status PASS", "PASS"),
        ("TC-07", "Order Complete", "Complete order after QC inspection", "Complete ORD-0001", "Product stock updated by +95 units", "PASS"),
        ("TC-08", "Shipment", "Create shipment within available stock", "Qty 95 finished goods", "Shipment created, stock reduced to 0", "PASS"),
        ("TC-09", "Shipment", "Create shipment exceeding available stock", "Qty 10 when stock is 0", "Insufficient stock error flash message", "PASS"),
        ("TC-10", "Export", "Export Inventory PDF Report", "GET /reports/inventory/pdf", "Valid 200 OK PDF binary download", "PASS")
    ]

    for tcid, mod, scen, inp, exp, stat in test_cases_data:
        row = t_test.add_row()
        row.cells[0].paragraphs[0].text = tcid
        row.cells[1].paragraphs[0].text = mod
        row.cells[2].paragraphs[0].text = scen
        row.cells[3].paragraphs[0].text = inp
        row.cells[4].paragraphs[0].text = exp
        row.cells[5].paragraphs[0].text = stat
        for c in row.cells:
            set_cell_margins(c)

    # CONCLUSION
    add_h1("CONCLUSION")
    add_p("The Vibe Clothing Company Smart Production & Manufacturing Management System successfully addresses the operational challenges of garment manufacturing facilities. By unifying raw material inventory tracking, Bill of Materials calculation, multi-stage batch advancement, quality control pass-rate analytics, outbound logistics dispatches, executive BI analytics, and activity audit trails into a single relational database platform, the application delivers complete operational transparency.")
    add_p("Strict adherence to business rules—including Indian Rupee (₹) financial formatting, Product ID identification (PROD-XXXX), server-side RBAC, automated stock ledgers, and multi-format report exports—ensures that the system meets enterprise standards for reliability, data integrity, and usability.")

    # BIBLIOGRAPHY
    add_h1("BIBLIOGRAPHY")
    add_bullet("Flask Documentation (v3.0.x): https://flask.palletsprojects.com/")
    add_bullet("SQLAlchemy & Flask-SQLAlchemy Documentation: https://flask-sqlalchemy.palletsprojects.com/")
    add_bullet("Python Software Foundation (v3.14): https://docs.python.org/3/")
    add_bullet("MySQL 8.0 Reference Manual: https://dev.mysql.com/doc/refman/8.0/en/")
    add_bullet("Bootstrap 5.3 Documentation: https://getbootstrap.com/docs/5.3/")
    add_bullet("Chart.js Data Visualization Library (v4.4): https://www.chartjs.org/docs/latest/")
    add_bullet("ReportLab PDF Generation Library: https://www.reportlab.com/documentation/")
    add_bullet("openpyxl Spreadsheet Engine: https://openpyxl.readthedocs.io/")

    # APPENDICES
    add_h1("APPENDICES")

    add_h2("D. Data Flow Diagram (DFD)")
    add_h3("Level 0 Context DFD:")
    add_p("[User / Admin / Manager / Staff] <---> (Vibe Clothing ERP System) <---> [MySQL / SQLite Database]")

    add_h3("Level 1 Process DFD:")
    add_bullet("1.0 Authentication Process: Validates user credentials against users table and creates session.")
    add_bullet("2.0 Inventory Process: Manages material master records, reorder alerts, and inventory transactions.")
    add_bullet("3.0 Production Process: Validates BOM against inventory, starts production orders, deducts raw stock, generates batches.")
    add_bullet("4.0 Batch Stepper Process: Advances garment assembly through Cutting → Dyeing → Stitching → Finishing → Packing.")
    add_bullet("5.0 Quality Control Process: Evaluates batch garments, computes pass rates, updates usable produced output.")
    add_bullet("6.0 Logistics Process: Creates shipments, validates finished goods stock, updates delivery status.")
    add_bullet("7.0 Analytics & Reporting Process: Queries database ledgers to render Chart.js graphs and generate PDF/Excel exports.")

    add_h2("E. Table Structure")

    db_schema_details = [
        ("Table: users", [
            ("id", "INTEGER", "PK, AUTO_INCREMENT", "Surrogate Primary Key"),
            ("user_id", "VARCHAR(20)", "UNIQUE", "Prefixed User ID (USR-0001)"),
            ("username", "VARCHAR(50)", "UNIQUE", "Login Username"),
            ("email", "VARCHAR(100)", "UNIQUE", "User Email Address"),
            ("password_hash", "VARCHAR(255)", "NOT NULL", "Werkzeug Scrypt Password Hash"),
            ("role", "VARCHAR(20)", "DEFAULT 'staff'", "RBAC Role (admin, manager, staff)"),
            ("is_active", "BOOLEAN", "DEFAULT TRUE", "Account Active Status")
        ]),
        ("Table: materials", [
            ("id", "INTEGER", "PK, AUTO_INCREMENT", "Surrogate Primary Key"),
            ("material_id", "VARCHAR(20)", "UNIQUE", "Prefixed Material ID (MAT-0001)"),
            ("name", "VARCHAR(100)", "NOT NULL", "Raw Material Name"),
            ("quantity", "FLOAT", "DEFAULT 0.0", "Current Available Stock Quantity"),
            ("unit", "VARCHAR(20)", "NOT NULL", "Unit of Measure (Meter, PCS, Liter)"),
            ("unit_cost", "FLOAT", "DEFAULT 0.0", "Cost per Unit in Indian Rupees (₹)"),
            ("minimum_level", "FLOAT", "DEFAULT 0.0", "Minimum Stock Level Threshold"),
            ("critical_level", "FLOAT", "DEFAULT 0.0", "Critical Stock Level Threshold"),
            ("supplier_id", "VARCHAR(20)", "FK -> suppliers", "Supplier Reference ID")
        ]),
        ("Table: products", [
            ("id", "INTEGER", "PK, AUTO_INCREMENT", "Surrogate Primary Key"),
            ("product_id", "VARCHAR(20)", "UNIQUE", "Prefixed Product ID (PROD-0001)"),
            ("name", "VARCHAR(100)", "NOT NULL", "Garment Product Name"),
            ("selling_price", "FLOAT", "NOT NULL", "Selling Price in Indian Rupees (₹)"),
            ("cost_price", "FLOAT", "NOT NULL", "Manufacturing Cost Price in ₹"),
            ("target_quantity", "INTEGER", "DEFAULT 0", "Target Production Quantity"),
            ("available_stock", "INTEGER", "DEFAULT 0", "Available Finished Goods Stock"),
            ("status", "VARCHAR(20)", "DEFAULT 'Active'", "Product Status (Active, In Production)")
        ])
    ]

    for tname, cols in db_schema_details:
        add_h3(tname)
        t_col = doc.add_table(rows=1, cols=4)
        t_col.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_table_header(t_col.rows[0], [Inches(1.5), Inches(1.2), Inches(1.8), Inches(2.0)], ["Field Name", "Data Type", "Constraints", "Description"])
        for fn, dt, cs, ds in cols:
            r = t_col.add_row()
            r.cells[0].paragraphs[0].text = fn
            r.cells[1].paragraphs[0].text = dt
            r.cells[2].paragraphs[0].text = cs
            r.cells[3].paragraphs[0].text = ds
            for c in r.cells:
                set_cell_margins(c)

    add_h2("F. Sample Coding")
    add_h3("1. ERPEngine Start Production Logic (services/erp_engine.py):")
    add_p("```python\n@staticmethod\ndef start_production(order_id, user_id='System'):\n    order = ProductionOrder.query.filter_by(order_id=order_id).first()\n    analysis = ERPEngine.check_material_availability(order.product_id, order.quantity)\n    if analysis['readiness_status'] == 'BLOCKED':\n        return False, 'Material Shortage! Cannot start production.'\n    for item in order.product.bom_items:\n        material = Material.query.filter_by(material_id=item.material_id).first()\n        qty_needed = item.quantity_required * order.quantity\n        material.quantity = max(0.0, material.quantity - qty_needed)\n    order.status = 'IN_PROGRESS'\n    db.session.commit()\n    return True, 'Production started successfully!'\n```")

    add_h3("2. Currency Formatter Utility (utils/helpers.py):")
    add_p("```python\ndef format_inr(value):\n    try:\n        val = float(value)\n        s, *d = f'{val:.2f}'.split('.')\n        r = s[-3:]\n        o = s[:-3]\n        if o:\n            r = ','.join([o[max(i-2,0):i] for i in range(len(o), 0, -2)][::-1]) + ',' + r\n        return f'₹{r}.{d[0]}'\n    except Exception:\n        return '₹0.00'\n```")

    add_h2("G. Sample Input Screens")
    add_bullet("Figure G.1: Login Screen (Full-Screen Video Background with Centered Auth Card)")
    add_bullet("Figure G.2: Add Raw Material Modal Input Form (Material Code MAT-XXXX, Unit Cost ₹)")
    add_bullet("Figure G.3: Bill of Materials (BOM) Editor Form (Product ID PROD-XXXX Specification)")
    add_bullet("Figure G.4: Quality Inspection Entry Form (Total Evaluated, Defective Units, Defect Type)")
    add_bullet("Figure G.5: Outbound Shipment Dispatch Form (Carrier, Tracking Number, Finished Goods Qty)")

    add_h2("H. Sample Output Screens")
    add_bullet("Figure H.1: Production Overview Dashboard (10 KPI Summary Cards & Donut Chart)")
    add_bullet("Figure H.2: Inventory Master Table (Status Badges: OK, LOW, CRITICAL & Valuation ₹)")
    add_bullet("Figure H.3: Production Kanban Board (PLANNED, IN PROGRESS, COMPLETED)")
    add_bullet("Figure H.4: Batch Visual Stage Stepper (Cutting → Dyeing → Stitching → Finishing → Packing)")
    add_bullet("Figure H.5: Quality Control Inspection Register (Pass/Fail Badges & Pass Rate %)")
    add_bullet("Figure H.6: Factory Analytics BI Dashboard (6 Chart.js Interactive Graphs)")
    add_bullet("Figure H.7: End-to-End Traceability Lineage Timeline Graph")

    # Save to desktop and brain directory
    desktop_docx_path = "/Users/vishnupriyachinnasamy24/Desktop/Vibe_Clothing_ERP_Phase_2_Documentation.docx"
    brain_docx_path = "/Users/vishnupriyachinnasamy24/.gemini/antigravity/brain/e58e4d33-f607-4084-a8b7-848caba7a7aa/Vibe_Clothing_ERP_Phase_2_Documentation.docx"

    doc.save(desktop_docx_path)
    doc.save(brain_docx_path)
    print(f"Documentation DOCX generated successfully at: {desktop_docx_path}")

if __name__ == "__main__":
    build_docx_documentation()
