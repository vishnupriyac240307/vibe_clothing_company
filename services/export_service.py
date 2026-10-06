import os
import io
import csv
from datetime import datetime, timezone

from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExportService:

    @staticmethod
    def export_pdf(title, headers, rows, summary_info=None, charts_flowables=None):
        """Generates a professional PDF report for Vibe Clothing Company with Logo, font-safe Rupee formatting, and visual charts."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Custom Styles
        title_style = ParagraphStyle(
            'CompanyTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#0f172a')
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=17,
            textColor=colors.HexColor('#2563eb')
        )
        meta_style = ParagraphStyle(
            'MetaStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            textColor=colors.HexColor('#64748b')
        )
        section_heading_style = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading3'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#0f172a'),
            spaceBefore=10,
            spaceAfter=6
        )
        header_cell_style = ParagraphStyle(
            'HeaderCell',
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.white,
            alignment=1
        )
        body_cell_style = ParagraphStyle(
            'BodyCell',
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#1e293b')
        )

        elements = []

        # 1. Company Logo & Company Name Header
        logo_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'static', 'images', 'logo.png'))
        if not os.path.exists(logo_path):
            logo_path = os.path.abspath('static/images/logo.png')

        if os.path.exists(logo_path):
            img = Image(logo_path, width=42, height=42)
            title_para = Paragraph("VIBE CLOTHING COMPANY", title_style)
            header_table = Table([[img, title_para]], colWidths=[52, 468])
            header_table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('PADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(header_table)
        else:
            elements.append(Paragraph("VIBE CLOTHING COMPANY", title_style))

        # 2. Subtitle - Strictly WITHOUT "ERP"
        clean_title = title.replace(' ERP', '').replace('ERP ', '').replace('ERP', '').strip()
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(f"Smart Manufacturing System — {clean_title}", subtitle_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(f"Generated On: {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M:%S UTC')} | Confidential Internal Document", meta_style))

        elements.append(Spacer(1, 8))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceAfter=12))

        # 3. Optional Summary Box
        if summary_info:
            summary_data = []
            for k, v in summary_info.items():
                clean_k = str(k).replace('₹', 'Rs. ')
                clean_v = str(v).replace('₹', 'Rs. ')
                summary_data.append([
                    Paragraph(f"<b>{clean_k}:</b>", meta_style),
                    Paragraph(clean_v, meta_style)
                ])
            summary_table = Table(summary_data, colWidths=[150, 350])
            summary_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
                ('PADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(summary_table)
            elements.append(Spacer(1, 12))

        # 4. Optional Visual Charts
        if charts_flowables:
            elements.append(Paragraph("Visual Analytics & Key Performance Charts", section_heading_style))
            elements.extend(charts_flowables)
            elements.append(Spacer(1, 8))

        # 5. Main Data Table (Replacing Unicode ₹ with Rs. to eliminate black box glyph issues in ReportLab PDF)
        elements.append(Paragraph("Detailed Data Records", section_heading_style))
        table_data = []
        formatted_headers = [Paragraph(str(h).replace('₹', 'Rs. '), header_cell_style) for h in headers]
        table_data.append(formatted_headers)

        for row in rows:
            formatted_row = [Paragraph(str(cell).replace('₹', 'Rs. '), body_cell_style) for cell in row]
            table_data.append(formatted_row)

        num_cols = len(headers)
        avail_width = 520
        col_width = avail_width / num_cols if num_cols > 0 else 100

        data_table = Table(table_data, colWidths=[col_width] * num_cols, repeatRows=1)
        data_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('PADDING', (0, 0), (-1, -1), 5),
        ]))

        elements.append(data_table)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Vibe Clothing Company — End of Report", meta_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_analytics_charts_flowables(analytics_data):
        """Generates ReportLab visual chart flowables for Analytics PDF exports."""
        from reportlab.graphics.shapes import Drawing
        from reportlab.graphics.charts.piecharts import Pie
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Table, TableStyle, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        styles = getSampleStyleSheet()
        chart_title_style = ParagraphStyle(
            'ChartTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.HexColor('#0f172a'),
            alignment=0
        )

        flowables = []

        # Chart 1: Production Status Distribution (Pie Chart)
        status_labels = analytics_data.get('status_labels', [])
        status_values = analytics_data.get('status_values', [])
        
        d1 = Drawing(250, 130)
        if status_values and sum(status_values) > 0:
            pie = Pie()
            pie.x = 20
            pie.y = 10
            pie.width = 100
            pie.height = 100
            pie.data = status_values
            pie.labels = [f"{lbl} ({val})" for lbl, val in zip(status_labels, status_values)]
            pie.sideLabels = 0
            
            color_map = {
                'COMPLETED': colors.HexColor('#16a34a'),
                'IN_PROGRESS': colors.HexColor('#f59e0b'),
                'PLANNED': colors.HexColor('#64748b'),
                'CANCELLED': colors.HexColor('#ef4444')
            }
            for idx, lbl in enumerate(status_labels):
                pie.slices[idx].fillColor = color_map.get(lbl, colors.HexColor('#2563eb'))
            d1.add(pie)

        # Chart 2: Inventory Valuation by Category (Bar Chart)
        cat_labels = analytics_data.get('cat_labels', [])[:5]
        cat_values = analytics_data.get('cat_values', [])[:5]

        d2 = Drawing(250, 130)
        if cat_values:
            bc = VerticalBarChart()
            bc.x = 40
            bc.y = 18
            bc.height = 90
            bc.width = 190
            bc.data = [[val / 1000 for val in cat_values]]
            bc.categoryAxis.categoryNames = [c[:10] for c in cat_labels]
            bc.categoryAxis.labels.fontSize = 7
            bc.valueAxis.labels.fontSize = 7
            bc.bars[0].fillColor = colors.HexColor('#6366f1')
            d2.add(bc)

        # Chart 3: Top Products Target Qty vs Stock (Grouped Bar Chart)
        prod_names = analytics_data.get('prod_names', [])[:4]
        prod_targets = analytics_data.get('prod_targets', [])[:4]
        prod_stocks = analytics_data.get('prod_stocks', [])[:4]

        d3 = Drawing(250, 130)
        if prod_names:
            bc3 = VerticalBarChart()
            bc3.x = 35
            bc3.y = 18
            bc3.height = 90
            bc3.width = 195
            bc3.data = [prod_targets, prod_stocks]
            bc3.categoryAxis.categoryNames = [p[:10] for p in prod_names]
            bc3.categoryAxis.labels.fontSize = 7
            bc3.valueAxis.labels.fontSize = 7
            bc3.bars[0].fillColor = colors.HexColor('#2563eb')
            bc3.bars[1].fillColor = colors.HexColor('#10b981')
            d3.add(bc3)

        # Chart 4: Quality Defects Breakdown (Bar Chart)
        qc_labels = analytics_data.get('qc_labels', [])[:4]
        qc_values = analytics_data.get('qc_values', [])[:4]

        d4 = Drawing(250, 130)
        if qc_values:
            bc4 = VerticalBarChart()
            bc4.x = 35
            bc4.y = 18
            bc4.height = 90
            bc4.width = 195
            bc4.data = [qc_values]
            bc4.categoryAxis.categoryNames = [q[:10] for q in qc_labels]
            bc4.categoryAxis.labels.fontSize = 7
            bc4.valueAxis.labels.fontSize = 7
            bc4.bars[0].fillColor = colors.HexColor('#ef4444')
            d4.add(bc4)

        # Combine into 2x2 grid layout table
        chart_table_data = [
            [
                Paragraph("<b>Production Status Distribution</b>", chart_title_style),
                Paragraph("<b>Raw Inventory Asset Valuation (in '000 Rs.)</b>", chart_title_style)
            ],
            [d1, d2],
            [
                Paragraph("<b>Top Products Target vs Stock (Blue=Target, Green=Stock)</b>", chart_title_style),
                Paragraph("<b>Garment Quality Defect Categories (Units)</b>", chart_title_style)
            ],
            [d3, d4]
        ]

        chart_table = Table(chart_table_data, colWidths=[255, 255])
        chart_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#f1f5f9')),
            ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#f8fafc')),
            ('BACKGROUND', (0, 2), (1, 2), colors.HexColor('#f8fafc')),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))

        flowables.append(chart_table)
        flowables.append(Spacer(1, 10))

        return flowables

    @staticmethod
    def export_excel(title, headers, rows):
        """Generates an Excel spreadsheet for Vibe Clothing Company."""
        wb = openpyxl.Workbook()
        ws = wb.active
        clean_title = title.replace(' ERP', '').replace('ERP ', '').replace('ERP', '').strip()
        ws.title = clean_title[:31]

        # Styling Definitions
        title_font = Font(name='Calibri', size=16, bold=True, color='0F172A')
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title Row - Without ERP
        clean_headers = [str(h).replace('₹', '₹ ') for h in headers]
        ws.append([f"VIBE CLOTHING COMPANY — Smart Manufacturing {clean_title}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Report Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"])

        ws.append([]) # Empty row

        # Header Row
        ws.append(clean_headers)
        header_row_idx = 4
        for col_num in range(1, len(clean_headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Data Rows
        for row in rows:
            ws.append([str(c) if c is not None else '' for c in row])

        # Apply borders and auto column width
        for row in ws.iter_rows(min_row=header_row_idx, max_row=ws.max_row, min_col=1, max_col=len(clean_headers)):
            for cell in row:
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_csv(headers, rows):
        """Generates CSV format."""
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        clean_headers = [str(h).replace('₹', 'INR ') for h in headers]
        writer.writerow(clean_headers)
        for row in rows:
            clean_row = [str(c).replace('₹', 'INR ') if c is not None else '' for c in row]
            writer.writerow(clean_row)
        return buffer.getvalue().encode('utf-8')
